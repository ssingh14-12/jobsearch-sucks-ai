from flask import Blueprint, request, jsonify, send_file
from app.services.claude_client import build_resume
import io

resume_bp = Blueprint('resume', __name__)


@resume_bp.route('/resume/build', methods=['POST'])
def build():
    data = request.get_json()
    jd_text        = data.get('jd_text', '').strip()
    resume_version = data.get('resume_version', 'V1').strip()
    company_name   = data.get('company_name', '').strip()

    if not jd_text:
        return jsonify({"error": "Job description is required."}), 400

    if len(jd_text) < 50:
        return jsonify({"error": "JD is too short — paste the full description."}), 400

    try:
        result = build_resume(jd_text, resume_version, company_name)
        return jsonify(result)
    except Exception as e:
        return jsonify({"error": f"Resume build failed: {str(e)}"}), 500


@resume_bp.route('/resume/download', methods=['POST'])
def download():
    """Generate and return a .docx file for the resume."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return jsonify({"error": "python-docx not installed. Run: pip install python-docx"}), 500

    data   = request.get_json()
    resume = data.get('resume', {})

    if not resume:
        return jsonify({"error": "Resume data is required."}), 400

    doc = Document()

    # ── Page margins ──
    for section in doc.sections:
        section.top_margin    = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin   = Inches(0.75)
        section.right_margin  = Inches(0.75)

    header = resume.get('header', {})

    # ── Name ──
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name_para.add_run(header.get('name', 'Soumya Singh'))
    run.bold      = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0, 0, 0)

    # ── Tagline ──
    if header.get('tagline'):
        tag_para = doc.add_paragraph()
        tag_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tag_run = tag_para.add_run(header['tagline'])
        tag_run.font.size = Pt(10)
        tag_run.font.color.rgb = RGBColor(80, 80, 80)

    # ── Contact line ──
    parts = [
        header.get('email', ''),
        header.get('phone', ''),
        header.get('linkedin', ''),
        header.get('location', ''),
    ]
    contact_line = '  |  '.join(p for p in parts if p)
    if contact_line:
        c_para = doc.add_paragraph()
        c_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c_run = c_para.add_run(contact_line)
        c_run.font.size = Pt(9)
        c_run.font.color.rgb = RGBColor(60, 60, 60)

    def add_section_heading(text):
        p = doc.add_paragraph()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0, 0, 0)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after  = Pt(2)
        # Add a bottom border via paragraph XML
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), '000000')
        pBdr.append(bottom)
        pPr.append(pBdr)

    # ── Summary ──
    if resume.get('summary'):
        add_section_heading('Professional Summary')
        s_para = doc.add_paragraph(resume['summary'])
        s_para.runs[0].font.size = Pt(10)

    # ── Experience ──
    experience = resume.get('experience', [])
    if experience:
        add_section_heading('Experience')
        for job in experience:
            # Title + Company line
            job_para = doc.add_paragraph()
            title_run = job_para.add_run(job.get('title', ''))
            title_run.bold      = True
            title_run.font.size = Pt(10)
            company_run = job_para.add_run(f"  —  {job.get('company', '')}")
            company_run.font.size = Pt(10)
            company_run.font.color.rgb = RGBColor(60, 60, 60)
            job_para.paragraph_format.space_after = Pt(0)

            # Dates + Location
            meta_para = doc.add_paragraph()
            meta_run = meta_para.add_run(f"{job.get('dates', '')}  ·  {job.get('location', '')}")
            meta_run.font.size        = Pt(9)
            meta_run.italic           = True
            meta_run.font.color.rgb   = RGBColor(100, 100, 100)
            meta_para.paragraph_format.space_after = Pt(2)

            # Bullets
            for bullet in job.get('bullets', []):
                b_para = doc.add_paragraph(style='List Bullet')
                b_run  = b_para.add_run(bullet)
                b_run.font.size = Pt(10)
                b_para.paragraph_format.space_after = Pt(1)

    # ── Education ──
    education = resume.get('education', [])
    if education:
        add_section_heading('Education')
        for edu in education:
            e_para = doc.add_paragraph()
            degree_run = e_para.add_run(edu.get('degree', ''))
            degree_run.bold      = True
            degree_run.font.size = Pt(10)
            school_run = e_para.add_run(f"  —  {edu.get('school', '')}")
            school_run.font.size = Pt(10)
            school_run.font.color.rgb = RGBColor(60, 60, 60)
            e_para.paragraph_format.space_after = Pt(0)

            dates_para = doc.add_paragraph()
            dates_run  = dates_para.add_run(edu.get('dates', ''))
            dates_run.font.size      = Pt(9)
            dates_run.italic         = True
            dates_run.font.color.rgb = RGBColor(100, 100, 100)
            if edu.get('note'):
                dates_run2 = dates_para.add_run(f"  ·  {edu['note']}")
                dates_run2.font.size      = Pt(9)
                dates_run2.font.color.rgb = RGBColor(100, 100, 100)

    # ── Certifications ──
    certs = resume.get('certifications', [])
    if certs:
        add_section_heading('Certifications')
        for cert in certs:
            c_para = doc.add_paragraph(style='List Bullet')
            c_run  = c_para.add_run(cert)
            c_run.font.size = Pt(10)

    # ── Skills ──
    skills = resume.get('skills', {})
    if skills:
        add_section_heading('Skills')
        for category, items in skills.items():
            sk_para = doc.add_paragraph()
            cat_run = sk_para.add_run(f"{category}: ")
            cat_run.bold      = True
            cat_run.font.size = Pt(10)
            items_run = sk_para.add_run(', '.join(items))
            items_run.font.size = Pt(10)
            sk_para.paragraph_format.space_after = Pt(2)

    # ── Save to BytesIO and return ──
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    company = data.get('company', 'Resume')
    filename = f"Soumya_Singh_{company.replace(' ', '_')}_Resume.docx"

    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
